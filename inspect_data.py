import openpyxl
import pandas as pd
import zipfile
import os

excel_path = r"C:\Users\Asus\.gemini\antigravity\scratch\smart-farming-pakcoy-new\new Dataset_Pakcoy.xlsx"
wb = openpyxl.load_workbook(excel_path, read_only=True)
print("=== EXCEL SHEETS ===")
print("Sheet names:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    print(f"\n--- SHEET: {sheet_name} ---")
    df = pd.read_excel(excel_path, sheet_name=sheet_name)
    print("Shape:", df.shape)
    print("Columns:", df.columns.tolist())
    print("\nHead (5 rows):")
    print(df.head())
    print("\nInfo / Dtypes:")
    print(df.dtypes)
    print("\nNull counts:")
    print(df.isnull().sum())
    print("\nDescribe summary:")
    print(df.describe(include='all'))

zip_files = [
    r"C:\Users\Asus\.gemini\antigravity\scratch\smart-farming-pakcoy-new\s3Masa Pertumbuhan.zip",
    r"C:\Users\Asus\.gemini\antigravity\scratch\smart-farming-pakcoy-new\s3Mendekati Panen pokcoy.zip",
    r"C:\Users\Asus\.gemini\antigravity\scratch\smart-farming-pakcoy-new\s3siap panen pokcoy.zip"
]

print("\n=== ZIP FILES ANALYSIS ===")
for z_path in zip_files:
    fname = os.path.basename(z_path)
    print(f"\n--- ZIP: {fname} ---")
    with zipfile.ZipFile(z_path, 'r') as z:
        file_list = z.namelist()
        print(f"Total entries: {len(file_list)}")
        # filter images
        img_exts = ('.jpg', '.jpeg', '.png', '.bmp', '.webp')
        images = [f for f in file_list if f.lower().endswith(img_exts) and not f.startswith('__MACOSX')]
        dirs = [f for f in file_list if f.endswith('/')]
        print(f"Total image files: {len(images)}")
        print("Sample image file paths (up to 10):", images[:10])
        print("Subdirectories:", dirs[:10])
